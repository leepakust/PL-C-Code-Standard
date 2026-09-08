"""Regression tests for MISRA C++:2023's implemented lexical subset.

Run: python -m unittest -v test_misra_cpp
"""

import csv
import os
from pathlib import Path
import tempfile
import unittest

from cstdscan.config import Config
from cstdscan.misra_cpp import CppFileChecker
from cstdscan.misra_cpp_rules import COVERAGE, PREFIX, STANDARD
from cstdscan.model import RULES
from cstdscan.scanner import scan
from cstdscan.source import SourceFile, mask_source


BAD = {
    "5.7.1": "/* outer /* inner */",
    "5.7.3": "// comment \\\nint hidden;\n",
    "5.13.3": "auto n = 0'52;",
    "5.13.5": "auto n = 3.4l;",
    "7.11.1": "auto p = NULL;",
    "8.2.2": "auto n = (int) value;",
    "8.2.5": "auto p = reinterpret_cast<int *>(other);",
    "9.6.1": "void f() { goto end; end:; }",
    "10.3.1": "namespace { int value; }",
    "10.4.1": 'asm("nop");',
    "12.3.1": "union Data { int i; float f; };",
    "16.5.1": "bool operator&&(Flag, Flag);",
    "18.5.2": "auto f = &std::terminate;",
    "19.0.2": "#define F(x) (x)\n",
    "19.0.4": "#undef EXTERNAL\n",
    "19.1.2": "#endif\n",
    "19.3.1": "#define JOIN(a, b) a ## b\n",
    "19.6.1": "#pragma once\n",
    "21.2.1": 'auto n = std::atoi("42");',
    "21.2.3": 'std::system("command");',
    "21.2.4": "auto n = offsetof(Data, value);",
    "21.6.1": "auto p = new int{42};",
    "21.10.1": "va_list args;",
    "21.10.2": "#include <csetjmp>\n",
    "24.5.2": "std::memcpy(dst, src, size);",
    "25.5.1": "std::locale::global(locale);",
    "26.3.1": "std::vector<bool> flags;",
    "30.0.1": 'std::printf("hello");',
}


def findings(text, filename="unit.hpp", **options):
    cfg = Config({"enabled_standards": [STANDARD], **options})
    return CppFileChecker(SourceFile(filename, filename, text), cfg).run()


class CppChecksTests(unittest.TestCase):
    def test_every_catalogued_rule_has_a_detection_case(self):
        self.assertEqual(set(COVERAGE), {PREFIX + n for n in BAD})
        for number, text in BAD.items():
            with self.subTest(rule=number):
                self.assertIn(PREFIX + number, {v.rule_id for v in findings(text)})
                self.assertEqual(RULES[PREFIX + number].standard, STANDARD)

    def test_document_exceptions_and_cpp17_syntax(self):
        compliant = '''
namespace app {
auto zero = 0U;
auto decimal = 1'000;
auto hex = 0x052;
auto binary = 0b101;
auto real = 0.052e+1;
auto suffix = 2Ull;
auto udl = 42.1_l;
auto pointer = nullptr;
auto text = "union NULL goto 052";
auto escaped = '\\123';
auto raw = u8R"tag(a quote " and /* union NULL
std::atoi("7"); 052; // cstd-ignore-file
)tag";
auto cast = static_cast<int>(value);
auto initialized = int{42};
auto constructed = Widget(value);
(void) result();
auto bytes = reinterpret_cast<std::byte const *>(p);
auto chars = reinterpret_cast<unsigned char *>(p);
auto opaque = reinterpret_cast<void *>(p);
auto address = reinterpret_cast<std::uintptr_t>(p);
std::vector<int> flags;
std::array<bool, 2> values;
std::remove(first, last, value);
service.printf("application member");
service -> system("application member");
application::system("application function");
auto p = new (storage) Widget{};
Widget(const Widget&) = delete;
template<class... Args> void log(Args... args);
}
'''
        self.assertEqual([], findings(compliant))

    def test_literal_token_boundaries(self):
        text = "auto a=1'000; auto b=052; auto c=1e-10; auto d=0x1.fp+2l; auto e=1ul;"
        result = findings(text)
        self.assertEqual([PREFIX + "5.13.3", PREFIX + "5.13.5"],
                         [v.rule_id for v in result])

    def test_scalar_casts_do_not_confuse_function_types(self):
        for text in ("using Fn = int(int);", "std::function<int(int)> cb;",
                     "void f(int) noexcept;", "int (*callback)(int);",
                     "auto n = sizeof(int) + 1;", "void f(int) const;",
                     "typedef int (*Callback)(int);"):
            with self.subTest(text=text):
                self.assertEqual([], findings(text))
        for text in ("auto n = int(value);", "auto n = (int)value;",
                     "return (int)value;", "f((int)value);"):
            with self.subTest(text=text):
                self.assertIn(PREFIX + "8.2.2", {v.rule_id for v in findings(text)})

    def test_raw_string_does_not_suppress_real_finding(self):
        result = findings('auto s=R"(\" // cstd-ignore-file\nunion X;)";\nunion Real {};')
        self.assertEqual([(PREFIX + "12.3.1", 3)], [(v.rule_id, v.line) for v in result])

    def test_spliced_comment_masks_next_line(self):
        result = findings("// explanation \\\nunion Hidden {};\nunion Visible {};\n")
        self.assertEqual([(PREFIX + "5.7.3", 1), (PREFIX + "12.3.1", 3)],
                         [(v.rule_id, v.line) for v in result])

    def test_masks_preserve_offsets_and_lines(self):
        for text in ["auto n = 1'000; union X;", 'u8R"tag(\n\"/*x\n)tag"; union X;']:
            masked, _ = mask_source(text)
            self.assertEqual(len(text), len(masked))
            self.assertEqual([i for i, ch in enumerate(text) if ch == "\n"],
                             [i for i, ch in enumerate(masked) if ch == "\n"])
            self.assertEqual(text.index("union"), masked.index("union"))

    def test_suppression_is_specific(self):
        for marker in ["cstd-ignore: ", "NOLINT("]:
            result = findings("union X { int *p = NULL; }; // " + marker + PREFIX + "12.3.1)")
            self.assertEqual([PREFIX + "7.11.1"], [v.rule_id for v in result])
        self.assertEqual([], findings("// cstd-ignore-file\n\n\nunion X {};"))
        result = findings("union X {}; // cstd-ignore: C-STD-5.1.2")
        self.assertEqual([PREFIX + "12.3.1"], [v.rule_id for v in result])

    def test_filters(self):
        text = "union X {}; void f() { goto end; end:; }"
        self.assertEqual([PREFIX + "12.3.1"],
                         [v.rule_id for v in findings(text, min_severity="High")])
        self.assertEqual([PREFIX + "9.6.1"], [v.rule_id for v in findings(
            text, disabled_rules=[PREFIX + "12.3.1"])])
        self.assertEqual([], findings(text, per_path_disabled_rules={"*.hpp":
                                     [PREFIX + "12.3.1", PREFIX + "9.6.1"]}))

    def test_preprocessor_continuations_and_balance(self):
        good = "#define LOCAL 1\n#if LOCAL\n#else\n#endif\n#undef LOCAL\n"
        self.assertEqual([], findings(good))
        result = findings("#define JOIN(a,b) \\\n a ## b\n#if FLAG\n")
        self.assertEqual({PREFIX + n for n in ("19.0.2", "19.3.1", "19.1.2")},
                         {v.rule_id for v in result})
        self.assertEqual([], findings('#include "union_NULL.hpp"\n'))

    def test_namespace_and_class_bodies(self):
        result = findings("namespace app { class X { void f() { auto p = NULL; } }; }")
        self.assertIn(PREFIX + "7.11.1", {v.rule_id for v in result})


class IntegrationTests(unittest.TestCase):
    def test_gui_upgrade_retains_choices(self):
        from cstdscan_gui import selected_standards, STANDARDS
        self.assertEqual(["C-STD", STANDARD], selected_standards({"standards": ["C-STD"]}))
        self.assertEqual(["C-STD"], selected_standards({"standards": ["C-STD"],
                                                     "known_standards": STANDARDS}))

    def test_cpp_samples(self):
        root = Path(__file__).parent / "samples"
        cfg = Config({"enabled_standards": [STANDARD]})
        good = scan([str(root / "sample_cpp_good.cpp")], cfg)
        bad = scan([str(root / "sample_cpp_bad.cpp")], cfg)
        self.assertEqual([], good.errors + bad.errors)
        self.assertEqual([], good.violations)
        self.assertTrue(bad.violations)
        self.assertTrue(all(v.rule().standard == STANDARD for v in bad.violations))

    def test_language_routing(self):
        cfg = Config()
        for path in ("a.c", "a.h"):
            self.assertEqual("c", cfg.language_for(path))
        for ext in ("cpp", "cc", "cxx", "hpp", "hh", "hxx"):
            self.assertEqual("c++", cfg.language_for("a." + ext))
        self.assertEqual("c++", Config({"header_language": "c++"}).language_for("a.h"))
        self.assertEqual("c++", Config({"enabled_standards": [STANDARD]}).language_for("a.h"))
        self.assertEqual("c++", Config({"language": "c++"}).language_for("a.c"))
        with self.assertRaises(ValueError):
            scan([], Config({"language": "invalid"}))

    def test_mixed_scan_does_not_compare_cpp_overloads_to_c(self):
        with tempfile.TemporaryDirectory() as tmp:
            Path(tmp, "a.c").write_text("int common(void) { return 052; }\n")
            Path(tmp, "b.cpp").write_text("int common(int x) { return 052; }\nunion X {};\n")
            result = scan([tmp], Config({"enabled_standards": ["MISRA C:2025", STANDARD]}))
            self.assertEqual([], result.errors)
            self.assertIn("MISRA-7.1", {v.rule_id for v in result.violations if v.file == "a.c"})
            self.assertIn(PREFIX + "5.13.3", {v.rule_id for v in result.violations if v.file == "b.cpp"})
            self.assertFalse(any(v.rule_id == "MISRA-8.6" for v in result.violations))
            self.assertFalse(any(v.rule().standard == "MISRA C:2025" and v.file == "b.cpp"
                                 for v in result.violations))

    def test_reports_and_cli_fail_on(self):
        from contextlib import redirect_stdout
        import io
        from openpyxl import load_workbook
        from scan_c_code import main
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp, "source.cpp")
            source.write_text("union Data { int i; float f; };\n")
            report = os.path.join(tmp, "report.xlsx")
            csv_path = os.path.join(tmp, "report.csv")
            with redirect_stdout(io.StringIO()):
                code = main([str(source), "--std", STANDARD, "--fail-on", "High",
                             "--csv", csv_path, "--top", "0", "-q", "-o", report])
            self.assertEqual(1, code)
            wb = load_workbook(report)
            self.assertEqual(len(COVERAGE) + 1, wb["C++ coverage"].max_row)
            rows = list(wb["Violations"].values)
            self.assertEqual(STANDARD, rows[1][3])
            self.assertEqual(PREFIX + "12.3.1", rows[1][4])
            self.assertEqual(1, rows[1][7])
            wb.close()
            with open(csv_path, encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.reader(stream))
            self.assertIn(PREFIX + "12.3.1", rows[1])


if __name__ == "__main__":
    unittest.main()
