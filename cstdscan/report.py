"""Excel and CSV report writers."""

import csv
import datetime
import os
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .model import RULES, SEVERITY_ORDER, rule_title
from .misra_cpp_rules import COVERAGE


HEADER_FILL = PatternFill("solid", fgColor="1F3B57")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=16, color="1F3B57")
SUB_FONT = Font(bold=True, size=11, color="1F3B57")
MONO_FONT = Font(name="Consolas", size=9)
THIN = Side(style="thin", color="D0D7DE")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SEVERITY_FILL = {
    "Critical": PatternFill("solid", fgColor="F8CBCB"),
    "High": PatternFill("solid", fgColor="FBE0C4"),
    "Medium": PatternFill("solid", fgColor="FCF3C6"),
    "Low": PatternFill("solid", fgColor="E7ECF0"),
}
SEVERITY_FONT = {
    "Critical": Font(color="8B0000", bold=True, size=10),
    "High": Font(color="9C4E00", bold=True, size=10),
    "Medium": Font(color="7A5C00", size=10),
    "Low": Font(color="4A5560", size=10),
}

VIOLATION_COLUMNS = [
    ("#", 6),
    ("Severity", 10),
    ("Class", 11),
    ("Standard", 20),
    ("Rule", 23),
    ("Rule title", 42),
    ("File", 38),
    ("Line", 7),
    ("Col", 6),
    ("Function", 22),
    ("Offending code", 60),
    ("What was found", 46),
    ("Why it matters", 52),
    ("How to fix", 48),
    ("Also cites", 20),
    ("Confidence", 11),
]


def _style_header(ws, row, count):
    for col in range(1, count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left",
                                   wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[row].height = 28


def _write_block(ws, row, col, title, headers, rows, widths=None):
    ws.cell(row=row, column=col, value=title).font = SUB_FONT
    row += 1
    for i, head in enumerate(headers):
        cell = ws.cell(row=row, column=col + i, value=head)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BOX
    header_row = row
    for r, values in enumerate(rows, start=1):
        for i, value in enumerate(values):
            cell = ws.cell(row=header_row + r, column=col + i, value=value)
            cell.border = BOX
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            if i == 0 and value in SEVERITY_FILL:
                cell.fill = SEVERITY_FILL[value]
                cell.font = SEVERITY_FONT[value]
    if widths:
        for i, width in enumerate(widths):
            letter = get_column_letter(col + i)
            current = ws.column_dimensions[letter].width or 0
            ws.column_dimensions[letter].width = max(current, width)
    return header_row + len(rows) + 2


def write_excel(path, violations, stats, config):
    wb = Workbook()

    _sheet_summary(wb.active, violations, stats, config)
    _sheet_violations(wb.create_sheet("Violations"), violations)
    _sheet_by_file(wb.create_sheet("By file"), violations, stats)
    _sheet_by_rule(wb.create_sheet("By rule"), violations)
    _sheet_rules(wb.create_sheet("Rules reference"), config)
    coverage = wb.create_sheet("C++ coverage")
    coverage.append(["Rule", "Implemented lexical coverage"])
    for rule_id, scope in COVERAGE.items():
        coverage.append([rule_id, scope])
    _style_header(coverage, 1, 2)
    coverage.column_dimensions["A"].width = 25
    coverage.column_dimensions["B"].width = 110
    coverage.freeze_panes = "A2"
    for row in coverage.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        coverage.row_dimensions[row[0].row].height = 32

    directory = os.path.dirname(os.path.abspath(path))
    if directory and not os.path.isdir(directory):
        os.makedirs(directory)
    wb.save(path)
    return path


# --------------------------------------------------------------------------
def _sheet_summary(ws, violations, stats, config):
    ws.title = "Summary"
    ws["A1"] = "C / C++ Coding Standard Scan"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ("Static scan against the house C coding standard, "
                "MISRA C:2025, MISRA C++:2023 and CWE. Lexical subset; not a compliance certification.")
    ws["A2"].font = Font(italic=True, color="4A5560")

    meta = [
        ("Scan root", stats["root"]),
        ("Generated", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Scanner version", stats["version"]),
        ("Files scanned", stats["file_count"]),
        ("Lines scanned", stats["line_count"]),
        ("Rules active", stats["active_rules"]),
        ("Violations found", len(violations)),
        ("Language", config["language"]),
        (".h language", config["header_language"]),
    ]
    row = 4
    for label, value in meta:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

    sev_counts = Counter(v.rule().severity for v in violations)
    row = _write_block(
        ws, row + 1, 1, "Violations by severity",
        ["Severity", "Count"],
        [[s, sev_counts.get(s, 0)]
         for s in sorted(SEVERITY_ORDER, key=SEVERITY_ORDER.get)],
        widths=[18, 10])

    std_counts = Counter(v.rule().standard for v in violations)
    row = _write_block(
        ws, row, 1, "Violations by standard",
        ["Standard", "Count"],
        [[s, c] for s, c in sorted(std_counts.items(),
                                   key=lambda kv: -kv[1])],
        widths=[18, 10])

    cls_counts = Counter(v.rule().rule_class for v in violations)
    row = _write_block(
        ws, row, 1, "Violations by rule class",
        ["Class", "Count"],
        [[s, c] for s, c in sorted(cls_counts.items(),
                                   key=lambda kv: -kv[1])],
        widths=[18, 10])

    cat_counts = Counter(v.rule().category for v in violations)
    _write_block(
        ws, 10, 4, "Violations by area",
        ["Area", "Count"],
        [[s, c] for s, c in sorted(cat_counts.items(),
                                   key=lambda kv: -kv[1])],
        widths=[22, 10])

    rule_counts = Counter(v.rule_id for v in violations)
    top_rules = [[rid, RULES[rid].severity, RULES[rid].title, count]
                 for rid, count in rule_counts.most_common(15)]
    _write_block(ws, 10, 7, "Most frequent rules",
                 ["Rule", "Severity", "Title", "Count"], top_rules,
                 widths=[16, 11, 52, 8])

    file_counts = Counter(v.file for v in violations)
    top_files = [[f, c] for f, c in file_counts.most_common(15)]
    _write_block(ws, 32, 7, "Files with the most violations",
                 ["File", "Count"], top_files, widths=[46, 8])

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 46
    ws.sheet_view.showGridLines = False


# --------------------------------------------------------------------------
def _sheet_violations(ws, violations):
    headers = [h for h, _w in VIOLATION_COLUMNS]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    for index, v in enumerate(violations, start=1):
        rule = v.rule()
        also = ", ".join(v.also)
        ws.append([
            index,
            rule.severity,
            rule.rule_class,
            rule.standard,
            v.rule_id,
            rule.title,
            v.file,
            v.line,
            v.column,
            v.function,
            v.code,
            v.detail,
            rule.why,
            rule.fix,
            also,
            v.confidence,
        ])
        row = ws.max_row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = BOX
            cell.alignment = Alignment(vertical="top", wrap_text=(col >= 6))
        sev_cell = ws.cell(row=row, column=2)
        sev_cell.fill = SEVERITY_FILL.get(rule.severity, SEVERITY_FILL["Low"])
        sev_cell.font = SEVERITY_FONT.get(rule.severity,
                                          SEVERITY_FONT["Low"])
        sev_cell.alignment = Alignment(vertical="top", horizontal="center")
        ws.cell(row=row, column=11).font = MONO_FONT

    for i, (_h, width) in enumerate(VIOLATION_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"
    if violations:
        ref = "A1:%s%d" % (get_column_letter(len(headers)), ws.max_row)
        table = Table(displayName="Violations", ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleLight1", showRowStripes=True)
        try:
            ws.add_table(table)
        except ValueError:                          # pragma: no cover
            ws.auto_filter.ref = ref
    else:
        ws.auto_filter.ref = "A1:%s1" % get_column_letter(len(headers))


# --------------------------------------------------------------------------
def _sheet_by_file(ws, violations, stats):
    headers = ["File", "Lines", "Total", "Critical", "High", "Medium", "Low",
               "Top rule in this file"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    per_file = defaultdict(list)
    for v in violations:
        per_file[v.file].append(v)

    rows = []
    for rel, lines in sorted(stats["file_lines"].items()):
        items = per_file.get(rel, [])
        counts = Counter(v.rule().severity for v in items)
        top = Counter(v.rule_id for v in items).most_common(1)
        rows.append([
            rel, lines, len(items),
            counts.get("Critical", 0), counts.get("High", 0),
            counts.get("Medium", 0), counts.get("Low", 0),
            "%s (%d)" % (top[0][0], top[0][1]) if top else "",
        ])
    rows.sort(key=lambda r: (-r[3], -r[4], -r[2], r[0]))
    for values in rows:
        ws.append(values)
        for col in range(1, len(headers) + 1):
            ws.cell(row=ws.max_row, column=col).border = BOX
        if values[3]:
            ws.cell(row=ws.max_row, column=4).fill = SEVERITY_FILL["Critical"]
        if values[4]:
            ws.cell(row=ws.max_row, column=5).fill = SEVERITY_FILL["High"]

    widths = [50, 8, 8, 9, 8, 9, 7, 30]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)),
                                      max(1, ws.max_row))


# --------------------------------------------------------------------------
def _sheet_by_rule(ws, violations):
    headers = ["Rule", "Standard", "Severity", "Class", "Title", "Count",
               "Files affected"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    counts = Counter(v.rule_id for v in violations)
    files = defaultdict(set)
    for v in violations:
        files[v.rule_id].add(v.file)

    for rid, count in sorted(
            counts.items(),
            key=lambda kv: (SEVERITY_ORDER[RULES[kv[0]].severity], -kv[1])):
        rule = RULES[rid]
        ws.append([rid, rule.standard, rule.severity, rule.rule_class,
                   rule.title, count, len(files[rid])])
        row = ws.max_row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = BOX
            cell.alignment = Alignment(vertical="top", wrap_text=(col == 5))
        ws.cell(row=row, column=3).fill = SEVERITY_FILL[rule.severity]

    widths = [25, 22, 10, 11, 60, 8, 14]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)),
                                      max(1, ws.max_row))


# --------------------------------------------------------------------------
def _sheet_rules(ws, config):
    headers = ["Rule", "Standard", "Area", "Severity", "Class", "Enabled",
               "Title", "Why it matters", "How to fix"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    for rid, rule in sorted(
            RULES.items(),
            key=lambda kv: (kv[1].standard,
                            SEVERITY_ORDER[kv[1].severity], kv[0])):
        enabled = "yes" if config.rule_enabled(rule) else "no"
        ws.append([rid, rule.standard, rule.category, rule.severity,
                   rule.rule_class, enabled, rule.title, rule.why, rule.fix])
        row = ws.max_row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = BOX
            cell.alignment = Alignment(vertical="top", wrap_text=(col >= 7))
        ws.cell(row=row, column=4).fill = SEVERITY_FILL[rule.severity]

    widths = [25, 22, 16, 10, 11, 9, 46, 62, 52]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)),
                                      max(1, ws.max_row))


# --------------------------------------------------------------------------
def write_csv(path, violations):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow([h for h, _w in VIOLATION_COLUMNS])
        for index, v in enumerate(violations, start=1):
            rule = v.rule()
            writer.writerow([index, rule.severity, rule.rule_class,
                             rule.standard, v.rule_id, rule.title, v.file,
                             v.line, v.column, v.function, v.code, v.detail,
                             rule.why, rule.fix, ", ".join(v.also),
                             v.confidence])
    return path


def cross_reference_titles(violations):
    """Titles for every cross-referenced rule seen, for console output."""
    out = {}
    for v in violations:
        for rid in v.also:
            out[rid] = rule_title(rid)
    return out
